"""PowerPoint Analysis Deck Generator using python-pptx and matplotlib charts."""

import io
import os
from datetime import datetime, timezone
from pathlib import Path
import matplotlib
matplotlib.use("Agg")  # Non-interactive backend
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from app.config import settings


def _create_sales_trend_chart(daily_trends: list[dict]) -> io.BytesIO:
    """Render daily sales trend chart to image buffer."""
    fig, ax = plt.subplots(figsize=(6.5, 3.8), dpi=200)
    fig.patch.set_facecolor("#ffffff")
    ax.set_facecolor("#f8fafc")

    if daily_trends:
        dates = [d["date"][-5:] for d in daily_trends]
        revenues = [d["revenue"] for d in daily_trends]
        ax.plot(dates, revenues, marker="o", color="#0f766e", linewidth=2.5, markersize=7)
        ax.fill_between(dates, revenues, color="#ccfbf1", alpha=0.5)

        for i, (x, y) in enumerate(zip(dates, revenues)):
            ax.annotate(f"₹{int(y)}", (x, y), textcoords="offset points", xytext=(0, 8), ha="center", fontsize=8, weight="bold", color="#1e293b")
    else:
        ax.text(0.5, 0.5, "No sales recorded in period", ha="center", va="center", color="#64748b")

    ax.set_title("Daily Sales Revenue (₹)", fontsize=11, weight="bold", color="#0f172a", pad=12)
    ax.grid(True, linestyle="--", alpha=0.5, color="#cbd5e1")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#94a3b8")
    ax.spines["bottom"].set_color("#94a3b8")
    plt.tight_layout()

    buf = io.BytesIO()
    plt.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _create_top_products_chart(top_products: list[dict]) -> io.BytesIO:
    """Render top products horizontal bar chart to image buffer."""
    fig, ax = plt.subplots(figsize=(6.5, 3.8), dpi=200)
    fig.patch.set_facecolor("#ffffff")
    ax.set_facecolor("#f8fafc")

    if top_products:
        names = [p["name"][:18] for p in reversed(top_products)]
        revenues = [p["revenue"] for p in reversed(top_products)]
        bars = ax.barh(names, revenues, color="#3b82f6", height=0.55, edgecolor="#1d4ed8")

        for bar in bars:
            width = bar.get_width()
            ax.text(width + max(revenues) * 0.02, bar.get_y() + bar.get_height() / 2, f"₹{int(width)}", va="center", ha="left", fontsize=8, weight="bold", color="#1e293b")
    else:
        ax.text(0.5, 0.5, "No product sales data", ha="center", va="center", color="#64748b")

    ax.set_title("Top 5 SKUs by Sales Revenue (₹)", fontsize=11, weight="bold", color="#0f172a", pad=12)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_visible(False)
    ax.spines["bottom"].set_color("#94a3b8")
    ax.xaxis.grid(True, linestyle="--", alpha=0.5, color="#cbd5e1")
    plt.tight_layout()

    buf = io.BytesIO()
    plt.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _create_payment_pie_chart(payment_methods: dict) -> io.BytesIO:
    """Render payment mode distribution pie chart."""
    fig, ax = plt.subplots(figsize=(5.5, 3.8), dpi=200)
    fig.patch.set_facecolor("#ffffff")

    labels = []
    sizes = []
    colors_list = ["#10b981", "#6366f1", "#f59e0b", "#ec4899"]

    for k, v in payment_methods.items():
        if v > 0:
            labels.append(k)
            sizes.append(v)

    if sizes:
        wedges, texts, autotexts = ax.pie(
            sizes,
            labels=labels,
            autopct="%1.1f%%",
            startangle=140,
            colors=colors_list[:len(sizes)],
            wedgeprops={"edgecolor": "white", "linewidth": 2},
        )
        for t in texts:
            t.set_fontsize(9)
            t.set_color("#1e293b")
        for at in autotexts:
            at.set_fontsize(8)
            at.set_weight("bold")
            at.set_color("white")
    else:
        ax.text(0.5, 0.5, "No payments recorded", ha="center", va="center", color="#64748b")

    ax.set_title("Payment Mode Share", fontsize=11, weight="bold", color="#0f172a", pad=12)
    plt.tight_layout()

    buf = io.BytesIO()
    plt.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def generate_analysis_deck(
    sales_data: dict,
    low_stock_items: list[dict] | None = None,
    shop_info: dict | None = None,
    output_path: str | Path | None = None,
) -> str:
    """
    Generate a 5-slide Executive PowerPoint presentation deck with embedded Matplotlib chart images.
    """
    shop_name = (shop_info and shop_info.get("shop_name")) or settings.DEFAULT_SHOP_NAME
    date_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    if output_path is None:
        target_file = settings.get_report_path(f"sales_analysis_{date_str}.pptx")
    else:
        target_file = Path(output_path)
        target_file.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)  # 16:9 widescreen layout
    blank_layout = prs.slide_layouts[6]

    # --- SLIDE 1: Title Slide ---
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(5.625))  # 1 = MSO_SHAPE.RECTANGLE
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(15, 23, 42)  # Dark slate
    bg1.line.color.rgb = RGBColor(15, 23, 42)

    tb1 = slide1.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(2.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = shop_name
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(45, 212, 191)  # Emerald teal

    p2 = tf1.add_paragraph()
    p2.text = f"Weekly Store Performance & Sales Analysis Deck"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(241, 245, 249)

    p3 = tf1.add_paragraph()
    p3.text = f"Period: Past {sales_data.get('period_days', 7)} Days  |  Generated on: {date_str} by Supermarket Ops Agent"
    p3.font.size = Pt(11)
    p3.font.color.rgb = RGBColor(148, 163, 184)

    # --- SLIDE 2: Executive Summary KPI Cards ---
    slide2 = prs.slides.add_slide(blank_layout)
    
    # Title
    t_box = slide2.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    t_p = t_box.text_frame.paragraphs[0]
    t_p.text = "Executive Summary — Store KPIs"
    t_p.font.size = Pt(20)
    t_p.font.bold = True
    t_p.font.color.rgb = RGBColor(15, 23, 42)

    kpis = [
        ("Total Sales Revenue", f"₹{sales_data.get('total_revenue', 0.0):.2f}", RGBColor(15, 118, 110)),
        ("Total Invoices Cut", f"{sales_data.get('total_bills', 0)}", RGBColor(59, 130, 246)),
        ("Average Ticket Size", f"₹{(sales_data.get('total_revenue', 0)/(sales_data.get('total_bills', 1) or 1)):.2f}", RGBColor(99, 102, 241)),
        ("Low Stock Alerts", f"{len(low_stock_items or [])} SKUs", RGBColor(239, 68, 68)),
    ]

    for i, (label, val, col) in enumerate(kpis):
        left = Inches(0.8 + i * 2.15)
        card = slide2.shapes.add_shape(1, left, Inches(1.5), Inches(2.0), Inches(2.8))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(248, 250, 252)
        card.line.color.rgb = RGBColor(226, 232, 240)

        ctb = slide2.shapes.add_textbox(left + Inches(0.1), Inches(1.8), Inches(1.8), Inches(2.0))
        ctf = ctb.text_frame
        ctf.word_wrap = True

        cp1 = ctf.paragraphs[0]
        cp1.text = label
        cp1.font.size = Pt(11)
        cp1.font.color.rgb = RGBColor(100, 116, 139)

        cp2 = ctf.add_paragraph()
        cp2.text = val
        cp2.font.size = Pt(22)
        cp2.font.bold = True
        cp2.font.color.rgb = col

    # --- SLIDE 3: Sales Trend Timeline ---
    slide3 = prs.slides.add_slide(blank_layout)
    t_box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    t_p3 = t_box3.text_frame.paragraphs[0]
    t_p3.text = "Daily Revenue Trends"
    t_p3.font.size = Pt(20)
    t_p3.font.bold = True
    t_p3.font.color.rgb = RGBColor(15, 23, 42)

    trend_buf = _create_sales_trend_chart(sales_data.get("daily_trends", []))
    slide3.shapes.add_picture(trend_buf, Inches(0.8), Inches(1.3), width=Inches(8.4), height=Inches(3.8))

    # --- SLIDE 4: Top SKUs by Revenue ---
    slide4 = prs.slides.add_slide(blank_layout)
    t_box4 = slide4.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    t_p4 = t_box4.text_frame.paragraphs[0]
    t_p4.text = "Top Selling Products & Velocity"
    t_p4.font.size = Pt(20)
    t_p4.font.bold = True
    t_p4.font.color.rgb = RGBColor(15, 23, 42)

    prod_buf = _create_top_products_chart(sales_data.get("top_products", []))
    slide4.shapes.add_picture(prod_buf, Inches(0.8), Inches(1.3), width=Inches(8.4), height=Inches(3.8))

    # --- SLIDE 5: Payment Method & Low Stock Health ---
    slide5 = prs.slides.add_slide(blank_layout)
    t_box5 = slide5.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    t_p5 = t_box5.text_frame.paragraphs[0]
    t_p5.text = "Payment Shares & Inventory Health"
    t_p5.font.size = Pt(20)
    t_p5.font.bold = True
    t_p5.font.color.rgb = RGBColor(15, 23, 42)

    pay_buf = _create_payment_pie_chart(sales_data.get("payment_methods", {}))
    slide5.shapes.add_picture(pay_buf, Inches(0.8), Inches(1.3), width=Inches(4.5), height=Inches(3.8))

    # Low stock list box on right side
    ltb = slide5.shapes.add_textbox(Inches(5.5), Inches(1.3), Inches(3.7), Inches(3.8))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    lp1 = ltf.paragraphs[0]
    lp1.text = "⚠️ Low Stock Items to Reorder:"
    lp1.font.size = Pt(12)
    lp1.font.bold = True
    lp1.font.color.rgb = RGBColor(220, 38, 38)

    if low_stock_items:
        for itm in low_stock_items[:5]:
            lp = ltf.add_paragraph()
            lp.text = f"• {itm['name']}: {itm['stock_qty']} left (reorder at {itm['reorder_level']})"
            lp.font.size = Pt(10)
            lp.font.color.rgb = RGBColor(30, 41, 59)
    else:
        lp = ltf.add_paragraph()
        lp.text = "✅ All SKUs are above reorder threshold."
        lp.font.size = Pt(10)
        lp.font.color.rgb = RGBColor(16, 185, 129)

    prs.save(str(target_file))
    return str(target_file)
